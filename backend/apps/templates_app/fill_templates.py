"""Prepare and inspect DOCX templates without generating or rewriting prose."""
import io
import re
import zipfile

from docx import Document
from docx.text.paragraph import Paragraph
from docxcompose.composer import Composer
from docxtpl import DocxTemplate
from jinja2 import nodes
from jinja2.sandbox import SandboxedEnvironment

from .content_library import full_template_path
from .fill_paths import canonical_path, path_parts
from .ingestion import discover_blocks, bind_alternatives
from .jinja_filters import JINJA_FILTERS
from .placeholders import convert_paragraph
from .template_variables import extract_template_variables_from_text, normalize_docxtpl_blocks
from .word_templates import block_template_path, style_template_path

MAX_DOCX_BYTES = 15 * 1024 * 1024
MAX_XML_BYTES = 40 * 1024 * 1024


class FillEnvironment(SandboxedEnvironment):
    def is_safe_callable(self, obj):
        return any(obj is function for function in self.globals.values())


def fill_environment(**kwargs):
    env = FillEnvironment(autoescape=True, **kwargs)
    env.globals.clear()
    env.globals.update(JINJA_FILTERS)
    env.filters.update(JINJA_FILTERS)
    return env


def validate_docx(content):
    if len(content) > MAX_DOCX_BYTES:
        raise ValueError("Please upload a DOCX smaller than 15 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_XML_BYTES:
                raise ValueError("The expanded DOCX is too large (40 MB maximum).")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("The file is not a Word DOCX document.")
    except zipfile.BadZipFile:
        raise ValueError("The file is not a valid DOCX document.") from None



def story_paragraphs(document):
    """All body/header/footer paragraphs, including nested tables and text boxes."""
    for part in document.part.package.parts:
        name = str(part.partname)
        if name == "/word/document.xml" or name.startswith(("/word/header", "/word/footer")):
            element = getattr(part, "element", None)
            if element is not None:
                for paragraph in element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                    yield Paragraph(paragraph, document)

def prepare_upload(content):
    validate_docx(content)
    document = Document(io.BytesIO(content))
    blocks = discover_blocks(document)
    original_paragraphs = list(document.paragraphs)
    converted = set()
    for index, paragraph in enumerate(story_paragraphs(document)):
        conversion = convert_paragraph(paragraph, f"placeholder_{index + 1}")
        converted.update(conversion.fields)
    choices = bind_alternatives(document, blocks, original_paragraphs)
    # Uploaded templates may carry the original author's identity in properties.
    for key in ("author", "last_modified_by", "comments", "keywords", "subject", "identifier"):
        setattr(document.core_properties, key, "")
    output = io.BytesIO()
    document.save(output)
    content = output.getvalue()
    fields = inspect_docx(content)
    return content, {"convertedFields": sorted(converted), "choices": choices, "fields": fields, "ai": "No AI used"}



def wrap_optional(document, block, *, full=False):
    expression = '{%p if fill_options["' + block.key + '"] %}'
    if not full:
        first = document.element.body[0]
        opening = document.add_paragraph(expression)._p
        first.addprevious(opening)
        document.add_paragraph("{%p endif %}")
        return
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text
        references_block = any(token in text for token in (f'blocks["{block.key}"]', f"blocks['{block.key}']", f"blocks.{block.key}."))
        if not references_block or not re.search(r"\{%p?\s+(if|for)\b", text):
            continue
        depth = 0
        for last in paragraphs[index:]:
            for tag in re.findall(r"\{%p?\s+(if|for|endif|endfor)\b", last.text):
                depth += 1 if tag in {"if", "for"} else -1
            if depth == 0:
                opening = document.add_paragraph(expression)._p
                closing = document.add_paragraph("{%p endif %}")._p
                paragraph._p.addprevious(opening)
                last._p.addnext(closing)
                return
    raise ValueError(f"Optional section '{block.label}' has no identifiable boundary. Prepare its section boundaries before filling this template.")

def library_docx(template):
    full = full_template_path(template)
    if full:
        document = Document(full)
        for block in template.blocks.filter(required=False):
            wrap_optional(document, block, full=True)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    style = style_template_path(template)
    master = Document(style) if style else Document()
    if style:
        for child in list(master.element.body):
            if not child.tag.endswith("}sectPr"):
                master.element.body.remove(child)
    composer = Composer(master) if style else None
    for block in template.blocks.all():
        path = block_template_path(template, block)
        if path:
            child = Document(path)
        else:
            child = Document()
            for paragraph in (block.body or "").split("\n"):
                child.add_paragraph(paragraph)
        if not block.required:
            wrap_optional(child, block)
        if composer is None:
            composer = Composer(child)
        else:
            composer.append(child)
    output = io.BytesIO()
    (composer or Composer(master)).save(output)
    return output.getvalue()


def template_sources(content):
    doc = DocxTemplate(io.BytesIO(content))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for name in archive.namelist():
            if name == "word/document.xml" or (name.startswith(("word/header", "word/footer")) and name.endswith(".xml")):
                xml = archive.read(name).decode("utf-8")
                if re.search(r"\{\{\s*(?:r|p)\s+", xml):
                    raise ValueError("Rich-text and subdocument expressions are not supported in Fill template. Use ordinary {{ field }} expressions.")
                yield normalize_docxtpl_blocks(doc.patch_xml(xml))


def inspect_docx(content):
    """Inventory fields; collection members are edited as one structured input."""
    env = fill_environment()
    paths, controls, collections = set(), set(), set()
    for source in template_sources(content):
        ast = env.parse(source)
        for node in ast.find_all((nodes.Include, nodes.Import, nodes.FromImport, nodes.Extends, nodes.Macro)):
            raise ValueError(f"Unsupported template instruction: {type(node).__name__}")
        for call in ast.find_all(nodes.Call):
            if not isinstance(call.node, nodes.Name) or call.node.name not in JINJA_FILTERS:
                raise ValueError("Unsupported Docassemble function or method. Use a field, loop, choice, or supported formatting filter.")
        for filter_node in ast.find_all(nodes.Filter):
            if filter_node.name not in env.filters or filter_node.name in {"safe", "attr"}:
                raise ValueError(f"Unsupported template filter: {filter_node.name}")
        for node in ast.find_all((nodes.Mul, nodes.Pow)):
            raise ValueError("Multiplication and exponentiation are not supported in uploaded templates.")
        for node in ast.find_all(nodes.Getitem):
            if not isinstance(node.arg, nodes.Const):
                raise ValueError("Use literal list indexes or a for-loop variable instead of a computed index.")
        paths.update(extract_template_variables_from_text(source))
        from .template_variables import TemplateVarVisitor

        class Controls(TemplateVarVisitor):
            def references(self, expression):
                visitor = TemplateVarVisitor()
                visitor._loop_stack = list(self._loop_stack)
                visitor.visit(expression)
                return visitor.results

            def visit_If(self, node):
                controls.update(self.references(node.test))
                self.generic_visit(node)

            def visit_For(self, node):
                references = self.references(node.iter)
                controls.update(references)
                collections.update(references)
                super().visit_For(node)

        Controls().visit(ast)
    # The existing visitor names loop member paths with [i]. The collection is
    # already offered as JSON; those members are not independent scalar fields.
    paths = {p for p in paths if not any(f"[{i}]" in p for i in ("i", "j", "k", "l", "m"))}
    collections = {p for p in collections if not any(f"[{i}]" in p for i in ("i", "j", "k", "l", "m"))}
    paths.update(collections)
    parsed = {}
    for path in sorted(paths):
        try:
            parts = path_parts(path)
        except ValueError as error:
            raise ValueError(f"Unsupported field {path}: {error}") from error
        if parts[0] in {"blocks", "document", "section", "loop"} or parts[0] in JINJA_FILTERS:
            continue
        # Remove prefix paths (e.g. 'clients' alongside 'clients[0].name').
        if path not in collections and any(other != path and (other.startswith(path + ".") or other.startswith(path + "[")) for other in paths):
            continue
        key = canonical_path(path)
        parsed[key] = {"path": path, "key": key, "label": path.removeprefix("fields.").replace("_", " "),
                       "kind": "json" if path in collections else "text", "required": path in controls}
    lines = [line for line in (context_segments(p.text, parsed) for p in story_paragraphs(Document(io.BytesIO(content)))) if line]
    for field in parsed.values():
        field["context"] = field_context(lines, field["key"])
    # In document order, so the form reads the way the document does; fields
    # used only in conditions or loops follow.
    def first_use(field):
        return next(((index, position) for index, line in enumerate(lines) for position, segment in enumerate(line)
                     if field["key"] in segment.get("fields", ())), (len(lines), 0))
    return sorted(parsed.values(), key=first_use)


EXPRESSION = re.compile(r"\{\{(.*?)\}\}")


def _slot_keys(expression, fields):
    keys = []
    for variable in extract_template_variables_from_text("{{" + expression + "}}"):
        try:
            key = canonical_path(variable)
        except ValueError:
            continue
        # A loop member ({{ tenants[0].name }}) belongs to its collection field.
        owner = next((k for k in fields if key == k or key.startswith(k + "[")), None)
        if owner and owner not in keys:
            keys.append(owner)
    return keys


def context_segments(text, fields):
    """One paragraph as text and field slots, so the screen can show which
    blank a field fills instead of a sentence of identical underscores."""
    text = re.sub(r"\{%.*?%\}", "", text).replace("\t", " ")
    segments, position = [], 0
    for match in EXPRESSION.finditer(text):
        segments.append({"text": text[position:match.start()]})
        keys = _slot_keys(match.group(1), fields)
        segments.append({"fields": keys} if keys else {"text": "…"})
        position = match.end()
    segments.append({"text": text[position:]})
    for segment in segments:
        if "text" in segment:
            segment["text"] = re.sub(r" {2,}", " ", segment["text"])
    if segments and "text" in segments[0]:
        segments[0]["text"] = segments[0]["text"].lstrip()
    if segments and "text" in segments[-1]:
        segments[-1]["text"] = segments[-1]["text"].rstrip()
    segments = [segment for segment in segments if segment.get("fields") or segment["text"]]
    return segments if any(segment.get("fields") or segment["text"].strip() for segment in segments) else []


def _has_words(line):
    return any(re.search(r"\w", segment.get("text", "")) for segment in line)


def field_context(lines, key, limit=500):
    """The first paragraph that uses the field, with a neighbouring line when
    the paragraph is only the blank (a signature line, say)."""
    for index, line in enumerate(lines):
        if not any(key in segment.get("fields", ()) for segment in line):
            continue
        context = list(line)
        if not _has_words(line):
            before = next((l for l in reversed(lines[:index]) if _has_words(l)), None)
            after = next((l for l in lines[index + 1:] if _has_words(l)), None)
            context = ([*before, {"break": True}] if before else []) + context + ([{"break": True}, *after] if after else [])
        total = 0
        for position, segment in enumerate(context):
            total += len(segment.get("text", ""))
            if total > limit:
                context = context[:position] + [{"text": "…"}]
                break
        return context
    return []
