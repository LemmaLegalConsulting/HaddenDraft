"""Building court-filing template packages from maintained YAML specs.

A motion and its memorandum have a shape no received Word original supplies
cleanly: a caption, the motion, a memorandum with a table of contents and a
table of authorities, numbered argument headings, an exhibit index, and a
certificate of service. The ingestion converter reads a Word original and
guesses that structure; here it is stated. Each spec in
``content/filing-templates/<slug>.yaml`` holds the wording and block
definitions, and this module writes the same package the converter writes --
``document-templates/<slug>/manifest.yaml`` and ``template.docx``, plus one
snippet per block under ``docx-snippets/<slug>/blocks/`` -- so drafting,
review, and export treat it like any other prepared template.

The packages are generated. Change the spec or this builder and run
``manage.py build_filing_templates``; never edit the output by hand.

The tables of contents and authorities are Word fields. The template carries
one TOA field per category from ``content/drafting-rules/table-of-authorities.yaml``,
and export marks the citations and fills them
(`apps.exporting.table_of_authorities`).

Output is byte-for-byte deterministic: rebuilding an unchanged spec leaves every
checksum alone, so the template index sees no change where there was none.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

import yaml
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from apps.exporting.authorities import load_config as load_authorities_config


SPEC_DIR = "filing-templates"
BUILDER_VERSION = "2026-09-26-filing-templates-v1"
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_CORE_TIME = datetime(2026, 1, 1)
TEXT_WIDTH_TWIPS = 9360  # 6.5 inches between one-inch margins

FIELD_REF_RE = re.compile(r"\bfields\.([a-z][a-z0-9_]*)")
ITALIC_RE = re.compile(r"\*([^*]+)\*")
LATITUDE_FILL_MODE = {
    "locked": "none",
    "guided": "constrained_generation",
    "generate": "constrained_generation",
}

TOC_PENDING_TEXT = "The table of contents is built when Word updates this document's fields."
TOA_PENDING_TEXT = "The table of authorities is built when Word updates this document's fields."


class FilingTemplateSpecError(ValueError):
    pass


def _checksum(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_spec(path: Path) -> dict:
    spec = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    required = {"schema_version", "slug", "title", "kind", "blocks", "document"}
    missing = sorted(required - set(spec))
    if missing:
        raise FilingTemplateSpecError(f"{path}: missing {', '.join(missing)}")
    if spec["slug"] != path.stem:
        raise FilingTemplateSpecError(f"{path}: slug {spec['slug']!r} must match the file name")
    keys = [block.get("key") for block in spec["blocks"]]
    if any(not key for key in keys) or len(keys) != len(set(keys)):
        raise FilingTemplateSpecError(f"{path}: block keys must be present and unique")
    placed = [item["block"] for item in spec["document"] if isinstance(item, dict) and "block" in item]
    unplaced = sorted(set(keys) - set(placed))
    unknown = sorted(set(placed) - set(keys))
    if unplaced or unknown:
        raise FilingTemplateSpecError(
            f"{path}: every block is placed once in the document (unplaced: {unplaced}, unknown: {unknown})"
        )
    for block in spec["blocks"]:
        if block.get("ai_latitude", "locked") not in LATITUDE_FILL_MODE:
            raise FilingTemplateSpecError(f"{path}: block {block['key']} has an unknown ai_latitude")
    for choice in spec.get("choices") or []:
        values = [option["value"] for option in choice.get("options") or []]
        if not values or choice.get("default") != values[0]:
            # An unanswered choice falls back to its default, and the default
            # must be the first alternative so no passage silently disappears.
            raise FilingTemplateSpecError(f"{path}: choice {choice.get('name')} must default to its first option")
    return spec


# --- Word document construction ---------------------------------------------


def _set_run_font(style, *, size=12, bold=None):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), "Times New Roman")


def _configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    _set_run_font(normal)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    body = styles.add_style("Brief Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    body.paragraph_format.first_line_indent = Inches(0.5)

    block_quote = styles.add_style("Brief Plain", WD_STYLE_TYPE.PARAGRAPH)
    block_quote.base_style = normal
    block_quote.paragraph_format.space_after = Pt(12)

    title = styles.add_style("Brief Title", WD_STYLE_TYPE.PARAGRAPH)
    title.base_style = normal
    _set_run_font(title, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    for level, alignment in ((1, WD_ALIGN_PARAGRAPH.CENTER), (2, WD_ALIGN_PARAGRAPH.LEFT)):
        heading = styles[f"Heading {level}"]
        _set_run_font(heading, bold=True)
        heading.font.italic = False
        heading.font.color.rgb = None
        heading.paragraph_format.alignment = alignment
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(12)
        heading.paragraph_format.keep_with_next = True
        # Theme fonts on the built-in headings would override Times New Roman.
        fonts = heading.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(attribute)) is not None:
                del fonts.attrib[qn(attribute)]

    toa_heading = styles.add_style("toa heading", WD_STYLE_TYPE.PARAGRAPH)
    toa_heading.element.set(qn("w:styleId"), "TOAHeading")
    toa_heading.base_style = normal
    _set_run_font(toa_heading, bold=True)
    toa_heading.paragraph_format.space_before = Pt(12)
    toa_heading.paragraph_format.space_after = Pt(6)
    toa_heading.paragraph_format.keep_with_next = True

    toa_entry = styles.add_style("table of authorities", WD_STYLE_TYPE.PARAGRAPH)
    toa_entry.element.set(qn("w:styleId"), "TableofAuthorities")
    toa_entry.base_style = normal
    toa_entry.paragraph_format.left_indent = Inches(0.25)
    toa_entry.paragraph_format.first_line_indent = Inches(-0.25)
    toa_entry.paragraph_format.space_after = Pt(6)
    toa_entry.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5), alignment=2, leader=1  # right-aligned, dot leader
    )


def _add_runs(paragraph, text, *, bold=False):
    """Add ``text`` as runs, italicizing ``*spans*``; Jinja passes through as text."""
    position = 0
    for match in ITALIC_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.bold = bold or None
        run = paragraph.add_run(match.group(1))
        run.italic = True
        run.bold = bold or None
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = bold or None


def _control(container, tag):
    """A paragraph holding one docxtpl paragraph-level tag, removed when rendered."""
    paragraph = container.add_paragraph(style="Brief Plain")
    paragraph.add_run(tag)
    return paragraph


def _paragraphs(text):
    return [part.strip() for part in re.split(r"\n\s*\n", str(text or "").strip()) if part.strip()]


def _block_style(block):
    return "Brief Plain" if block.get("type") in {"signature", "certificate"} or block.get("list") else "Brief Body"


def write_block(container, block):
    """Write one block's body the way a prepared template renders it.

    A generated block loops over the draft's paragraphs. A maintained block
    renders its maintained wording, with its formatting, unless the draft holds
    a revision, which then renders paragraph by paragraph.
    """
    key = block["key"]
    ref = f'blocks["{key}"]'
    style = _block_style(block)
    heading = block.get("heading_when_filled")
    if heading:
        _control(container, f'{{%p if {ref}["items"] %}}')
        container.add_paragraph(heading, style="Heading 1")
    if block.get("ai_latitude") == "generate" or block.get("list"):
        _control(container, f'{{%p for item in {ref}["items"] %}}')
        container.add_paragraph("{{ item }}", style=style)
        _control(container, "{%p endfor %}")
    else:
        _control(container, f'{{%p if {ref}["revision"] %}}')
        _control(container, f'{{%p for paragraph in {ref}["paragraphs"] %}}')
        container.add_paragraph("{{ paragraph }}", style=style)
        _control(container, "{%p endfor %}")
        _control(container, "{%p else %}")
        for text in _paragraphs(block.get("body")):
            if text.startswith("{%") and text.endswith("%}"):
                _control(container, text.replace("{%", "{%p", 1))
                continue
            paragraph = container.add_paragraph(style=style)
            for index, line in enumerate(text.split("\n")):
                if index:
                    paragraph.add_run().add_break()
                _add_runs(paragraph, line)
        _control(container, "{%p endif %}")
    if heading:
        _control(container, "{%p endif %}")


def _field_paragraph(document, instruction, pending_text, *, style="Brief Plain"):
    """A paragraph holding one complex field, marked dirty so Word rebuilds it."""
    paragraph = document.add_paragraph(style=style)

    def char(kind, *, dirty=False):
        run = OxmlElement("w:r")
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        if dirty:
            element.set(qn("w:dirty"), "true")
        run.append(element)
        paragraph._p.append(run)

    char("begin", dirty=True)
    run = OxmlElement("w:r")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    run.append(text)
    paragraph._p.append(run)
    char("separate")
    paragraph.add_run(pending_text).italic = True
    char("end")
    return paragraph


def _table_of_contents(document):
    document.add_paragraph("TABLE OF CONTENTS", style="Brief Title")
    _field_paragraph(document, 'TOC \\o "1-2" \\h \\z \\u', TOC_PENDING_TEXT)


def _table_of_authorities(document, authorities_config):
    document.add_paragraph("TABLE OF AUTHORITIES", style="Brief Title")
    passim = " \\p" if authorities_config.get("passim") else ""
    for category in authorities_config.get("categories") or []:
        document.add_paragraph(category["heading"], style="toa heading")
        _field_paragraph(
            document,
            f'TOA \\c "{int(category["number"])}"{passim}',
            TOA_PENDING_TEXT,
            style="table of authorities",
        )


def _set_cell_borders(cell, **edges):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        value = edges.get(edge)
        element.set(qn("w:val"), "single" if value else "nil")
        if value:
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), "000000")
        borders.append(element)
    properties.append(borders)


def _caption(document, title):
    court = document.add_paragraph(style="Brief Title")
    court.paragraph_format.space_before = Pt(0)
    court.add_run('IN THE {{ (court if "court" in (court|lower) else fields.court_name)|upper }}').bold = True
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(3.4)
    right.width = Inches(3.1)
    _set_cell_borders(left, right=True, bottom=True)
    _set_cell_borders(right)
    left.paragraphs[0].style = document.styles["Brief Plain"]
    for index, line in enumerate(["{{ plaintiff }}", "", "Plaintiff,", "", "v.", "", "{{ defendant }}", "", "Defendant."]):
        paragraph = left.paragraphs[0] if index == 0 else left.add_paragraph(style="Brief Plain")
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run(line)
    right.paragraphs[0].style = document.styles["Brief Plain"]
    for index, line in enumerate(["Case No. {{ case_number }}", "", "Magistrate {{ fields.magistrate }}", "", title]):
        paragraph = right.paragraphs[0] if index == 0 else right.add_paragraph(style="Brief Plain")
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Inches(0.2)
        run = paragraph.add_run(line)
        run.bold = line == title or None
    document.add_paragraph(style="Brief Plain")


def _page_break(document):
    document.add_paragraph(style="Brief Plain").add_run().add_break(WD_BREAK.PAGE)


def _new_document():
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
    _configure_styles(document)
    # The default template starts with one empty paragraph.
    body = document.element.body
    for paragraph in list(body.iterchildren(qn("w:p"))):
        body.remove(paragraph)
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.created = FIXED_CORE_TIME
    properties.modified = FIXED_CORE_TIME
    properties.revision = 1
    return document


def _add_page_numbers(document):
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        run = OxmlElement("w:r")
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        elif text == " PAGE ":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        else:
            element = OxmlElement("w:t")
            element.text = text
        run.append(element)
        paragraph._p.append(run)


def _deterministic_bytes(document) -> bytes:
    """Save ``document`` with fixed zip timestamps and member order."""
    raw = io.BytesIO()
    document.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist(), key=lambda item: (item != "[Content_Types].xml", item)):
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return output.getvalue()


def build_template_docx(spec, authorities_config) -> bytes:
    document = _new_document()
    blocks = {block["key"]: block for block in spec["blocks"]}
    for item in spec["document"]:
        if item == "page_break":
            _page_break(document)
        elif item == "table_of_contents":
            _table_of_contents(document)
        elif item == "table_of_authorities":
            _table_of_authorities(document, authorities_config)
        elif isinstance(item, dict) and "caption" in item:
            _caption(document, item["caption"])
        elif isinstance(item, dict) and "title" in item:
            document.add_paragraph(item["title"], style="Brief Title")
        elif isinstance(item, dict) and "heading" in item:
            document.add_paragraph(item["heading"], style=f"Heading {int(item.get('level', 1))}")
        elif isinstance(item, dict) and "block" in item:
            write_block(document, blocks[item["block"]])
        else:
            raise FilingTemplateSpecError(f"{spec['slug']}: unknown document item {item!r}")
    _add_page_numbers(document)
    return _deterministic_bytes(document)


def build_block_docx(block) -> bytes:
    document = _new_document()
    if block.get("label"):
        document.add_paragraph(block["label"], style="Heading 2")
    write_block(document, block)
    return _deterministic_bytes(document)


# --- Manifest ----------------------------------------------------------------


def _referenced_fields(spec):
    text = yaml.safe_dump(spec)
    fields = sorted(set(FIELD_REF_RE.findall(text)) | {"court_name", "magistrate"})
    return [f"fields.{name}" for name in fields]


def _block_manifest_row(spec, block, order, docx_path, checksum):
    latitude = block.get("ai_latitude", "locked")
    return {
        "key": block["key"],
        "label": block.get("label") or block["key"].replace("-", " ").title(),
        "type": block.get("type", "optional_clause"),
        "order": order,
        "required": bool(block.get("required", True)),
        "editable": bool(block.get("editable", latitude != "locked")),
        "ai_latitude": latitude,
        "ai_fill_mode": LATITUDE_FILL_MODE[latitude],
        "instructions": list(block.get("instructions") or []),
        "body": "\n\n".join(ITALIC_RE.sub(r"\1", text) for text in _paragraphs(block.get("body"))),
        "docx": docx_path,
        "sha256": checksum,
        "input": {"type": "list" if block.get("list") else "rich_text"},
        "lexical": {"node": "list" if block.get("list") else "paragraphs"},
    }


def build_package(spec_path: Path, content_root: Path, *, authorities_config=None) -> dict:
    """Write one template package under ``content_root``. Returns what changed."""
    spec = load_spec(spec_path)
    authorities_config = authorities_config or load_authorities_config()
    slug = spec["slug"]
    package_dir = content_root / "document-templates" / slug
    snippets_dir = content_root / "docx-snippets" / slug / "blocks"
    package_dir.mkdir(parents=True, exist_ok=True)
    snippets_dir.mkdir(parents=True, exist_ok=True)

    written = []

    def write(path: Path, data: bytes):
        if not path.exists() or path.read_bytes() != data:
            path.write_bytes(data)
            written.append(path)

    rows = []
    for order, block in enumerate(spec["blocks"], start=1):
        data = build_block_docx(block)
        write(snippets_dir / f"{block['key']}.docx", data)
        rows.append(
            _block_manifest_row(
                spec, block, order * 10, f"docx-snippets/{slug}/blocks/{block['key']}.docx", _checksum(data)
            )
        )
    write(package_dir / "template.docx", build_template_docx(spec, authorities_config))

    spec_bytes = spec_path.read_bytes()
    manifest = {
        "schema_version": 2,
        "slug": slug,
        "title": spec["title"],
        "kind": spec["kind"],
        "description": spec.get("description", ""),
        "goal": spec.get("goal", ""),
        "negative_goal": spec.get("negative_goal", ""),
        "aliases": list(spec.get("aliases") or []),
        "jurisdiction": spec.get("jurisdiction", ""),
        "source_label": spec.get("source_label", "Content library"),
        "active": bool(spec.get("active", True)),
        "verification": spec.get("verification", "unverified"),
        "render": {"strategy": "full_document", "docx": "template.docx"},
        "source": {
            "path": f"{SPEC_DIR}/{spec_path.name}",
            "sha256": _checksum(spec_bytes),
            "converter": "apps.templates_app.filing_templates",
            "converter_version": BUILDER_VERSION,
        },
        "responds_to": spec.get("responds_to") or {},
        "table_of_authorities": True,
        "fields": _referenced_fields(spec),
        "flags": [],
        "choices": list(spec.get("choices") or []),
        "blocks": rows,
    }
    manifest_bytes = (
        "# Generated by manage.py build_filing_templates from "
        f"{SPEC_DIR}/{spec_path.name}. Do not edit.\n"
        + yaml.safe_dump(manifest, sort_keys=False, allow_unicode=True, width=100)
    ).encode("utf-8")
    write(package_dir / "manifest.yaml", manifest_bytes)
    return {"slug": slug, "written": written}


def spec_paths(content_root: Path):
    return sorted((content_root / SPEC_DIR).glob("*.yaml"))


def build_all(content_root: Path) -> list[dict]:
    authorities_config = load_authorities_config()
    return [
        build_package(path, content_root, authorities_config=authorities_config)
        for path in spec_paths(content_root)
    ]
