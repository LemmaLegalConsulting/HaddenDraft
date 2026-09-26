import io
import tempfile
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase, override_settings
from docx import Document

from apps.drafting import opposing_filing
from apps.drafting.models import DraftingSession
from apps.drafting.services import create_draft
from apps.exporting.services import render_docx
from apps.matters.models import Matter
from apps.templates_app.content_library import load_manifest, sync_prepared_templates
from apps.templates_app.filing_templates import spec_paths
from apps.templates_app.models import DocumentTemplate


FILING_TEMPLATE_SLUGS = {
    "motion-for-summary-judgment",
    "motion-to-dismiss",
    "reply-in-support-of-motion-for-summary-judgment",
    "reply-in-support-of-motion-to-dismiss",
}


def _upload_docx(*paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class FilingTemplatePackageTests(TestCase):
    def test_every_spec_has_a_current_generated_package(self):
        # A spec edited without rebuilding its package would draft from stale wording.
        call_command("build_filing_templates", "--check", "--no-sync")
        self.assertEqual({path.stem for path in spec_paths(Path(settings.CONTENT_LIBRARY_DIR))}, FILING_TEMPLATE_SLUGS)

    def test_manifests_load_and_declare_what_they_answer(self):
        root = Path(settings.CONTENT_LIBRARY_DIR) / "document-templates"
        manifests = {slug: load_manifest(root / slug / "manifest.yaml")[0] for slug in FILING_TEMPLATE_SLUGS}
        self.assertTrue(manifests["reply-in-support-of-motion-to-dismiss"]["responds_to"]["required"])
        self.assertTrue(manifests["reply-in-support-of-motion-for-summary-judgment"]["responds_to"]["required"])
        self.assertFalse(manifests["motion-to-dismiss"]["responds_to"]["required"])
        self.assertEqual(manifests["motion-to-dismiss"]["responds_to"]["expects"][0], "complaint")
        for manifest in manifests.values():
            self.assertTrue(manifest["table_of_authorities"])
            for choice in manifest["choices"]:
                # An unanswered choice renders its first alternative, never nothing.
                self.assertEqual(choice["default"], choice["options"][0]["value"])


@override_settings(AI_DRAFTING_ENABLED=False)
class FilingTemplateDraftingTests(TestCase):
    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory()
        self.addCleanup(self.scratch.cleanup)
        # Only the public library: a developer's private templates must not
        # decide whether these pass.
        isolated = override_settings(
            ORGANIZATION_CONTENT_LIBRARY_DIR=str(Path(self.scratch.name) / "organization"),
            PUBLISHED_CONTENT_LIBRARY_DIR=str(Path(self.scratch.name) / "published"),
            DOCUMENT_STORAGE_BACKEND="filesystem",
            DOCUMENT_STORAGE_ROOT=str(Path(self.scratch.name) / "storage"),
        )
        isolated.enable()
        self.addCleanup(isolated.disable)
        sync_prepared_templates()
        self.user = User.objects.create_user("advocate", password="pw")
        self.matter = Matter.objects.create(
            external_id="LS-TOA-1",
            client_name="Jane Tenant",
            matter_type="Eviction defense",
            jurisdiction="Cleveland Municipal Court, Housing Division",
            source_system="Manual",
            summary="Landlord accepted rent after the termination date.",
        )

    def _session(self, slug, **template_data):
        template = DocumentTemplate.objects.get(slug=slug)
        return DraftingSession.objects.create(
            mode="draft_from_template",
            matter=self.matter,
            template=template,
            template_data={"court_case_number": "2026-CVG-000123", "plaintiff_name": "Acme Properties LLC", **template_data},
        )

    def _export_text(self, draft):
        content, reports = render_docx(draft)
        text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)
        return text, reports

    def test_sync_indexes_the_filing_templates_with_their_requirements(self):
        templates = {template.slug: template for template in DocumentTemplate.objects.filter(slug__in=FILING_TEMPLATE_SLUGS)}
        self.assertEqual(set(templates), FILING_TEMPLATE_SLUGS)
        reply = templates["reply-in-support-of-motion-for-summary-judgment"]
        self.assertTrue(reply.metadata["respondsTo"]["required"])
        self.assertTrue(reply.metadata["tableOfAuthorities"])
        self.assertEqual(reply.metadata["verification"], "starter")
        argument = reply.blocks.get(key="reply-argument")
        self.assertEqual(argument.ai_latitude, "generate")
        self.assertTrue(argument.ai_instructions)
        # The fallback marker is a Jinja literal, so sync never turned it into a question.
        self.assertNotIn("fields.attorney", argument.body)

    def test_motion_for_summary_judgment_exports_a_native_table_of_authorities(self):
        draft = create_draft(self._session("motion-for-summary-judgment"))
        text, reports = self._export_text(draft)
        self.assertIn("IN THE CLEVELAND MUNICIPAL COURT, HOUSING DIVISION", text)
        self.assertIn("DEFENDANT'S MOTION FOR SUMMARY JUDGMENT", text)
        self.assertIn("Now comes Defendant Jane Tenant", text)
        self.assertIn("CERTIFICATE OF SERVICE", text)
        self.assertNotIn("{{", text)
        report = reports["tableOfAuthorities"]
        entries = {item["longCite"] for item in report["authorities"]}
        self.assertIn("Temple v. Wean United, Inc., 50 Ohio St.2d 317, 364 N.E.2d 267 (1977)", entries)
        self.assertIn("Dresher v. Burt, 75 Ohio St.3d 280, 662 N.E.2d 264 (1996)", entries)
        self.assertIn("Civ.R. 56", entries)
        self.assertEqual(report["pageNumbers"], "unmeasured")
        # Without a model the generated sections say they need drafting; they are not blank.
        argument = next(section for section in draft.sections if section["key"] == "argument")
        self.assertIn("[Attorney review required: draft the argument]", argument["body"])

    def test_motion_to_dismiss_follows_the_chosen_ground_in_draft_and_export(self):
        default = create_draft(self._session("motion-to-dismiss"))
        motion = next(section for section in default.sections if section["key"] == "motion")
        self.assertIn("Civ.R. 12(B)(6)", motion["body"])
        self.assertNotIn("12(B)(1)", motion["body"])
        text, _reports = self._export_text(default)
        self.assertIn("O'Brien v. Univ. Community Tenants Union", text)
        self.assertNotIn("Spurlock", text)

        jurisdiction = create_draft(
            self._session("motion-to-dismiss", dismissal_ground="lack_of_subject_matter_jurisdiction")
        )
        text, reports = self._export_text(jurisdiction)
        self.assertIn("Civ.R. 12(B)(1)", text)
        self.assertIn("State ex rel. Bush v. Spurlock", text)
        self.assertNotIn("O'Brien", text)
        self.assertIn("Civ.R. 12", {item["longCite"] for item in reports["tableOfAuthorities"]["authorities"]})

    def test_a_reply_refuses_to_draft_until_the_opposition_is_identified(self):
        session = self._session("reply-in-support-of-motion-for-summary-judgment")
        with self.assertRaisesRegex(ValueError, "brief in opposition"):
            create_draft(session)

        opposing_filing.upload_filing(
            session,
            _upload_docx("PLAINTIFF'S BRIEF IN OPPOSITION", "Defendant's affidavit is self-serving."),
            filename="2026-09-12 Plaintiff's Brief in Opposition to MSJ.docx",
            user=self.user,
            filed_on=opposing_filing.parse_filed_on("2026-09-12"),
        )
        session.refresh_from_db()
        draft = create_draft(session)
        opening = next(section for section in draft.sections if section["key"] == "opening")
        self.assertIn("replies to Plaintiff's Brief in Opposition to MSJ, filed September 12, 2026", opening["body"])
        text, _reports = self._export_text(draft)
        self.assertIn("replies to Plaintiff's Brief in Opposition to MSJ, filed September 12, 2026", text)
