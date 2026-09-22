"""Versioned publication pipeline for operator-maintained legal sources.

The database is the publication pointer and searchable index; object storage is
the immutable artifact ledger.  A source becomes visible only after all raw,
validated, and published artifacts have been written successfully.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import mimetypes
import re
import tempfile
import threading
from datetime import datetime, timedelta
from pathlib import Path

from django.db import transaction
from django.db.models import Max
from django.utils import timezone
from docx import Document
from pypdf import PdfReader

from apps.core.storage import PUBLISHED, RAW, VALIDATED, copy_area, get_document_storage
from apps.sources.models import (
    ManagedSource, ManagedSourceChunk, ManagedSourceEvent, ManagedSourceVersion,
)

logger = logging.getLogger(__name__)

PARSER_VERSION = "managed-source-parser-v1"
CHUNKER_VERSION = "paragraph-chunker-v1"
MAX_CHUNK_CHARS = 6000
# A "validating" row older than this is treated as an interrupted worker, not
# as work in progress, so a re-upload can resume it.
STALE_VALIDATION = timedelta(minutes=30)
# Declared types that really are plain text. A browser labels RTF and HTML
# "text/..." too, and decoding those as text would index control words and
# markup as retrievable legal authority.
TEXT_CONTENT_TYPES = {"text/plain", "text/markdown", "text/x-markdown"}


class PublicationError(ValueError):
    pass


def _extension(filename):
    suffix = Path(filename or "source.txt").suffix.lower()
    return suffix if re.fullmatch(r"\.[a-z0-9]{1,10}", suffix or "") else ".bin"


def _extract(content, filename, content_type):
    """Return ``(text, pages)``; ``pages`` is ``None`` when nothing paginates.

    Only a PDF carries page boundaries the parser can read. A DOCX, Markdown,
    or text file has no page count until something renders it, so it reports
    pages as unmeasured rather than as zero.
    """
    suffix = Path(filename or "").suffix.lower()
    declared = (content_type or "").split(";")[0].strip().casefold()
    pages = None
    if declared == "application/pdf" or suffix == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(content))
            pages = [(page.extract_text() or "").replace("\x00", "").strip() for page in reader.pages]
        except Exception as exc:
            raise PublicationError(f"The PDF could not be read: {exc}") from exc
        text = "\n\n".join(page for page in pages if page)
    elif suffix == ".docx" or declared.endswith("wordprocessingml.document"):
        try:
            document = Document(io.BytesIO(content))
            text = "\n\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
        except Exception as exc:
            raise PublicationError(f"The DOCX could not be read: {exc}") from exc
    elif suffix in {".txt", ".md", ".text"} or declared in TEXT_CONTENT_TYPES:
        text = content.decode("utf-8-sig", "replace")
    else:
        raise PublicationError("Upload a PDF, DOCX, Markdown, or UTF-8 plain-text file.")
    text = text.replace("\x00", "").strip()
    if not text:
        raise PublicationError("No readable text was found. Upload OCR/plain text for a scanned document.")
    return text, pages


def _paragraphs(text, pages):
    """Paragraphs paired with the page they came from, ``None`` if unknown."""
    if pages is None:
        return [(part.strip(), None) for part in re.split(r"\n\s*\n", text) if part.strip()]
    return [
        (part.strip(), number)
        for number, page in enumerate(pages, 1)
        for part in re.split(r"\n\s*\n", page)
        if part.strip()
    ]


def _chunk_records(text, pages):
    """Pack paragraphs into chunks, carrying the page range of each one."""
    paragraphs = _paragraphs(text, pages)
    if not paragraphs:
        return []
    records, current = [], []
    first_page = last_page = None

    def flush():
        nonlocal current, first_page, last_page
        if current:
            records.append({
                "text": "\n\n".join(current), "page_start": first_page, "page_end": last_page,
            })
        current, first_page, last_page = [], None, None

    for paragraph, page in paragraphs:
        pieces = [paragraph[i:i + MAX_CHUNK_CHARS] for i in range(0, len(paragraph), MAX_CHUNK_CHARS)]
        for piece in pieces:
            size = sum(len(item) + 2 for item in current)
            if current and size + len(piece) > MAX_CHUNK_CHARS:
                flush()
            current.append(piece)
            if page is not None:
                first_page = page if first_page is None else min(first_page, page)
                last_page = page if last_page is None else max(last_page, page)
    flush()
    return records


def _proposed_case_metadata(text):
    """Conservative proposals for human review; never claims model verification."""
    lines = [line.strip() for line in text.splitlines() if line.strip()][:80]
    joined = "\n".join(lines)
    # Require text on both sides of the "v.", so a heading numbered with a
    # Roman numeral ("V. CONCLUSION") is not read as a case name.
    title = next((line for line in lines[:15] if re.search(r"\S\s+vs?\.?\s+\S", line, re.I)), "")
    citation_match = re.search(r"\b\d{4}-Ohio-\d+\b", joined)
    date_match = re.search(
        r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s+\d{1,2},\s+\d{4}\b", joined, re.I,
    )
    court = next((line for line in lines[:30] if "court" in line.casefold()), "")
    result = {}
    if title:
        result["title"] = title[:500]
    if citation_match:
        result["citation"] = citation_match.group(0)
    if court:
        result["court"] = court[:255]
    if date_match:
        try:
            result["decision_date"] = datetime.strptime(date_match.group(0), "%B %d, %Y").date()
        except ValueError:
            pass
    return result


def _manifest(version, chunks):
    source = version.source
    return {
        "schema_version": 1,
        "source": {"id": source.pk, "slug": source.slug, "kind": source.kind, "title": source.title},
        "version": {
            "id": version.pk, "number": version.number, "label": version.label,
            "source_sha256": version.sha256, "source_modified_at": (
                version.source_modified_at.isoformat() if version.source_modified_at else None
            ),
            "source_etag": version.source_etag, "imported_at": version.imported_at.isoformat(),
            "parser_version": version.parser_version, "chunker_version": version.chunker_version,
        },
        "metadata": {
            "jurisdiction": source.jurisdiction, "court": source.court, "county": source.county,
            "municipality": source.municipality, "appellate_district": source.appellate_district,
            "decision_date": source.decision_date.isoformat() if source.decision_date else None,
            "citation": source.citation, "publication_status": source.publication_status,
            "source_locator": source.source_locator, "source_system_id": source.source_system_id,
        },
        "chunks": [
            {
                "ordinal": index,
                "sha256": hashlib.sha256(record["text"].encode()).hexdigest(),
                "characters": len(record["text"]),
                "page_start": record["page_start"], "page_end": record["page_end"],
            }
            for index, record in enumerate(chunks, 1)
        ],
    }


def _record_failure(version, *, actor=None, error="", prefix=""):
    """Mark a version failed, and say so in its history.

    Every path that abandons an import goes through here: a version left in
    "uploaded" or "validating" with no error reads to an operator as work
    still in progress, which is the one thing a dead import must never look
    like.
    """
    message = f"{prefix}{error}"
    version.status = "failed"
    version.error = message
    version.validation_report = {"valid": False, "error": message}
    version.save(update_fields=["status", "error", "validation_report"])
    # The caller may hold a source instance from before a just-completed
    # publication. Read the pointer under the current database state before
    # deciding whether this failed candidate affects the source-level badge.
    current = ManagedSource.objects.get(pk=version.source_id)
    if not current.current_version_id:
        current.state = "failed"
        current.save(update_fields=["state", "updated_at"])
    ManagedSourceEvent.objects.create(
        source=current, version=version, action="failed", actor=actor, detail={"error": message},
    )
    return version


def is_resumable(version):
    """Whether re-importing the same bytes should retry this version.

    A row whose worker never started ("uploaded"), one killed mid-parse and
    stale enough that no worker is plausibly still on it, and one that
    recorded a failure are all retryable; anything validated is not.
    """
    if version.status in {"uploaded", "failed"}:
        return True
    if version.status != "validating":
        return False
    started = version.validation_started_at
    return started is None or timezone.now() - started > STALE_VALIDATION


def is_publishable(version):
    """Whether this version has validated artifacts to publish."""
    return bool(version.validated_manifest_key and version.validation_report.get("valid"))


def import_version(*, source, content, filename="source.txt", content_type="", label="", actor=None,
                   source_modified_at=None, source_etag="", publish=False, background=False):
    """Validate and stage an immutable version; optionally make it live."""
    if not isinstance(content, bytes):
        content = bytes(content)
    digest = hashlib.sha256(content).hexdigest()
    existing = source.versions.filter(sha256=digest).first()
    created = existing is None
    if existing:
        # A background import first creates the durable row and raw object, then
        # calls back through here in its worker to do the expensive parsing.
        # Re-uploading the same file is also how an operator retries a version
        # whose worker never finished.
        if is_resumable(existing):
            version = existing
        else:
            # Publishing a version that never validated raises; a re-upload is
            # a retry, not an error page, so leave the status to speak for it.
            if publish and source.current_version_id != existing.id and is_publishable(existing):
                publish_version(existing, actor=actor)
            return existing, False
    else:
        content_type = content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
        version = None
        with transaction.atomic():
            # (source, number) is unique, so two concurrent imports must not
            # read the same maximum: take the source row lock publication uses.
            locked = ManagedSource.objects.select_for_update().get(pk=source.pk)
            if not locked.versions.filter(sha256=digest).exists():
                number = (locked.versions.aggregate(value=Max("number"))["value"] or 0) + 1
                raw_key = (
                    f"managed-sources/{source.slug}/versions/{number}-{digest[:12]}"
                    f"/source{_extension(filename)}"
                )
                version = ManagedSourceVersion.objects.create(
                    source=source, number=number, label=label, original_filename=filename,
                    content_type=content_type, size_bytes=len(content), sha256=digest,
                    source_modified_at=source_modified_at, source_etag=source_etag, imported_by=actor,
                    parser_version=PARSER_VERSION, chunker_version=CHUNKER_VERSION, raw_key=raw_key,
                )
                ManagedSourceEvent.objects.create(
                    source=source, version=version, action="uploaded", actor=actor,
                )
        if version is None:
            # Another import recorded these exact bytes while we waited for the
            # lock; handle it the way a duplicate found earlier is handled.
            return import_version(
                source=source, content=content, filename=filename, content_type=content_type,
                label=label, actor=actor, source_modified_at=source_modified_at,
                source_etag=source_etag, publish=publish, background=background,
            )

    try:
        get_document_storage(RAW).put_bytes(
            content=content, key=version.raw_key, content_type=version.content_type,
        )
    except Exception as exc:
        _record_failure(version, actor=actor, error=exc, prefix="Raw storage failed: ")
        return version, created

    if background:
        source_id = source.pk
        actor_id = getattr(actor, "pk", None)
        version_id = version.pk

        def work():
            from django.contrib.auth import get_user_model
            from django.db import close_old_connections

            close_old_connections()
            worker_actor = None
            try:
                worker_actor = (
                    get_user_model().objects.filter(pk=actor_id).first() if actor_id else None
                )
                worker_source = ManagedSource.objects.get(pk=source_id)
                import_version(
                    source=worker_source, content=content, filename=filename, content_type=content_type,
                    label=label, actor=worker_actor, source_modified_at=source_modified_at,
                    source_etag=source_etag, publish=publish,
                )
            except Exception as exc:
                logger.exception("Background import of managed source version %s failed", version_id)
                try:
                    _record_failure(
                        ManagedSourceVersion.objects.get(pk=version_id), actor=worker_actor,
                        error=exc, prefix="Background validation failed: ",
                    )
                except Exception:
                    logger.exception(
                        "Could not record the failed background import of version %s", version_id,
                    )
            finally:
                close_old_connections()

        # Start the worker only once this row and its raw object are durable.
        # The thread has its own connection: inside the caller's transaction it
        # cannot see the version it was given, and on SQLite it deadlocks.
        transaction.on_commit(
            lambda: threading.Thread(
                target=work, name=f"managed-source-import-{version_id}", daemon=True,
            ).start()
        )
        return version, created

    try:
        version.status = "validating"
        version.error = ""
        version.validation_started_at = timezone.now()
        version.save(update_fields=["status", "error", "validation_started_at"])
        text, pages = _extract(content, filename, content_type)
        records = _chunk_records(text, pages)
        if not records:
            raise PublicationError("The extracted text did not produce any retrieval chunks.")
        proposed = _proposed_case_metadata(text) if source.kind == "case" else {}
        changed = []
        for field, value in proposed.items():
            current = getattr(source, field)
            # Only the literal placeholders are overwritten. A reviewed title
            # that happens to match the uploaded filename is still a title
            # somebody chose.
            replace_placeholder_title = (
                field == "title"
                and str(current).strip().casefold() in {"uploaded decision", "untitled decision"}
            )
            if not current or replace_placeholder_title:
                setattr(source, field, value)
                changed.append(field)
        if changed:
            source.save(update_fields=[*changed, "updated_at"])
        manifest = _manifest(version, records)
        prefix = f"managed-sources/{source.slug}/versions/{version.number}-{digest[:12]}"
        manifest_key = f"{prefix}/manifest.json"
        chunks_key = f"{prefix}/chunks.jsonl"
        encoded_manifest = json.dumps(manifest, indent=2, sort_keys=True).encode()
        encoded_chunks = b"".join(
            json.dumps({
                "ordinal": index, "text": record["text"],
                "page_start": record["page_start"], "page_end": record["page_end"],
            }, ensure_ascii=False).encode() + b"\n"
            for index, record in enumerate(records, 1)
        )
        validated = get_document_storage(VALIDATED)
        validated.put_bytes(content=encoded_manifest, key=manifest_key, content_type="application/json")
        validated.put_bytes(content=encoded_chunks, key=chunks_key, content_type="application/x-ndjson")
        with transaction.atomic():
            version.chunks.all().delete()
            ManagedSourceChunk.objects.bulk_create([
                ManagedSourceChunk(
                    version=version, ordinal=index, text=record["text"],
                    heading=(record["text"].splitlines()[0][:500] if record["text"] else ""),
                    page_start=record["page_start"], page_end=record["page_end"],
                    sha256=hashlib.sha256(record["text"].encode()).hexdigest(),
                ) for index, record in enumerate(records, 1)
            ])
            version.validated_manifest_key = manifest_key
            version.chunk_count = len(records)
            version.validation_report = {
                "valid": True, "characters": len(text),
                # A page count exists only where the format has one to read.
                # Reporting zero for a DOCX would state an unmeasured property
                # as a measurement.
                "pages": len(pages) if pages is not None else None,
                "pages_status": "measured" if pages is not None else "unmeasured",
                "proposed_metadata": {key: str(value) for key, value in proposed.items()},
            }
            version.status = "pending_review"
            version.save(update_fields=[
                "validated_manifest_key", "chunk_count", "validation_report", "status",
            ])
            ManagedSource.objects.filter(
                pk=source.pk, current_version__isnull=True, state="failed",
            ).update(state="draft", updated_at=timezone.now())
        ManagedSourceEvent.objects.create(
            source=source, version=version, action="validated", actor=actor,
            detail={"chunk_count": len(records)},
        )
        if publish:
            publish_version(version, actor=actor)
        return version, created
    except Exception as exc:
        _record_failure(version, actor=actor, error=exc)
        return version, created


def publish_version(version, *, actor=None):
    version = ManagedSourceVersion.objects.select_related("source").get(pk=version.pk)
    source = version.source
    if (
        not version.validated_manifest_key
        or not version.chunks.exists()
        or not version.validation_report.get("valid")
    ):
        raise PublicationError("Only a successfully validated version can be published.")
    validated = get_document_storage(VALIDATED)
    published = get_document_storage(PUBLISHED)
    prefix = version.validated_manifest_key.rsplit("/", 1)[0]
    keys = list(validated.iter_keys(prefix))
    if not keys or version.validated_manifest_key not in keys:
        raise PublicationError("Validated artifacts are missing; the live corpus was not changed.")
    try:
        # Stream through the shared copier rather than reading each object into
        # memory: a validated chunk file for a treatise runs to megabytes.
        copy_area(
            validated, published, prefix=prefix, skip_existing=False,
            content_type=lambda key: (
                "application/json" if key.endswith(".json") else "application/x-ndjson"
            ),
        )
        published_source_key = f"{prefix}/source{_extension(version.original_filename)}"
        with tempfile.TemporaryDirectory() as staging:
            staged = Path(staging) / f"source{_extension(version.original_filename)}"
            get_document_storage(RAW).download_to(version.raw_key, staged)
            published.put_file(
                local_path=staged, key=published_source_key, content_type=version.content_type,
            )
    except Exception as exc:
        version.error = f"Publication storage failed: {exc}"
        version.save(update_fields=["error"])
        raise PublicationError(version.error) from exc
    now = timezone.now()
    with transaction.atomic():
        locked = ManagedSource.objects.select_for_update().get(pk=source.pk)
        if locked.current_version_id and locked.current_version_id != version.id:
            ManagedSourceVersion.objects.filter(pk=locked.current_version_id).update(status="superseded")
        version.status = "published"
        version.published_at = now
        version.retired_at = None
        version.published_manifest_key = version.validated_manifest_key
        version.published_source_key = published_source_key
        version.error = ""
        version.save(update_fields=[
            "status", "published_at", "retired_at", "published_manifest_key", "published_source_key", "error",
        ])
        locked.current_version = version
        locked.state = "published"
        locked.last_successful_refresh_at = now
        locked.last_indexed_at = now
        locked.save(update_fields=[
            "current_version", "state", "last_successful_refresh_at", "last_indexed_at", "updated_at",
        ])
        ManagedSourceEvent.objects.create(source=locked, version=version, action="published", actor=actor)
    from apps.sources.research.index import reset_index
    reset_index()
    return version


def retire_source(source, *, actor=None):
    now = timezone.now()
    source.state = "retired"
    source.save(update_fields=["state", "updated_at"])
    if source.current_version_id:
        ManagedSourceVersion.objects.filter(pk=source.current_version_id).update(retired_at=now)
    ManagedSourceEvent.objects.create(
        source=source, version=source.current_version, action="retired", actor=actor,
    )
    from apps.sources.research.index import reset_index
    reset_index()


def rollback_source(source, version, *, actor=None):
    if version.source_id != source.id:
        raise PublicationError("That version belongs to a different source.")
    result = publish_version(version, actor=actor)
    ManagedSourceEvent.objects.create(source=source, version=version, action="rolled_back", actor=actor)
    return result
