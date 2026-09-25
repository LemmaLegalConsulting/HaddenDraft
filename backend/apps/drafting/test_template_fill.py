import io
import json
import tempfile
import zipfile
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document

from apps.drafting.models import DraftingSession, TemplateFillJob
from apps.exporting.template_fill import render_fill
from apps.templates_app.fill_paths import canonical_path, read_path
from apps.templates_app.fill_templates import inspect_docx, prepare_upload
from apps.templates_app.fill_mappings import seed_mappings, mapping_index, mapped_value
from apps.templates_app.models import DocumentTemplate, TemplateBlock, TemplateFieldMapping, FillTemplateUpload
from apps.matters.models import Matter


def docx_bytes(body, *, stories=False):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Maintained wording: ").bold = True
    paragraph.add_run(body)
    if stories:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "{{ fields.table_value }}"
        document.sections[0].header.paragraphs[0].text = "Header {{ fields.header_value }}"
        document.sections[0].footer.paragraphs[0].text = "Footer {{ fields.footer_value }}"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@override_settings(TEMPLATE_FILL_BACKGROUND=False, DOCUMENT_STORAGE_BACKEND="filesystem", AI_DRAFTING_ENABLED=True)
class TemplateFillTests(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.storage_settings = override_settings(DOCUMENT_STORAGE_ROOT=self.directory.name)
        self.storage_settings.enable()
        self.addCleanup(self.storage_settings.disable)
        self.user = get_user_model().objects.create_superuser("filler", "filler@example.test", "pw")
        self.client.force_login(self.user)
        self.matter = Matter.objects.create(external_id="fill-test", client_name="Alex Example", matter_type="Housing", jurisdiction="Ohio", raw_payload={"first": "Alex", "last": "Example", "client_full_name": "Alex Example", "client_address_home": {"city": "Cleveland"}})

    def upload(self, body, **kwargs):
        response = self.client.post("/api/template-fill/start/", {"matterId": self.matter.external_id, "file": SimpleUploadedFile("example.docx", docx_bytes(body, **kwargs))})
        self.assertEqual(response.status_code, 202, response.content)
        job = response.json()["job"]
        self.assertEqual(job["status"], "complete", job)
        return DraftingSession.objects.get(pk=job["sessionId"])

    def export(self, session, answers=None):
        return self.client.post(f"/api/template-fill/sessions/{session.pk}/", json.dumps({"answers": answers or {}, "revision": session.fill_state["revision"], "saveToLegalServer": False}), content_type="application/json")

    @patch("apps.ai.openai_client.OpenAICompatibleClient.complete", side_effect=AssertionError("No model calls allowed"))
    def test_upload_map_partial_export_and_provenance(self, complete):
        session = self.upload("{{ clients[0].name.first }} lives in {{ clients[0].address.city }}. Hearing: {{ fields.hearing_date }}", stories=True)
        state = session.fill_state
        first = next(f for f in state["fields"] if f["path"] == "clients[0].name.first")
        self.assertEqual(first["value"], "Alex")
        self.assertEqual(first["source"], "LegalServer: first")
        response = self.export(session)
        self.assertEqual(response.status_code, 202)
        job = response.json()["job"]
        self.assertEqual(job["status"], "complete", job)
        download = self.client.get(f"/api/template-fill/jobs/{job['id']}/file/")
        self.assertEqual(download.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(download.content)) as archive:
            for name in ("word/document.xml", "word/header1.xml", "word/footer1.xml"):
                text = archive.read(name).decode()
                self.assertIn('w:val="yellow"', text)
                self.assertNotIn("{{", text)
                self.assertNotIn("FILL", text)
            text = archive.read("word/document.xml").decode()
            self.assertIn("Alex", text)
            self.assertIn("Cleveland", text)
            self.assertIn("[Enter hearing date]", text)
            self.assertIn("<w:b", text)
        self.assertTrue(session.drafts.first().components.exists())
        complete.assert_not_called()

    def test_clear_mapped_value_snapshot_and_conflict(self):
        session = self.upload("{{ clients[0].name.first }}")
        key = session.fill_state["fields"][0]["key"]
        self.matter.raw_payload = {"first": "Changed later"}
        self.matter.save()
        url = f"/api/template-fill/sessions/{session.pk}/"
        response = self.client.patch(url, json.dumps({"answers": {key: ""}, "revision": 0}), content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["session"]["fields"][0]["value"], "")
        self.assertEqual(response.json()["session"]["fields"][0]["state"], "manual")
        conflict = self.client.patch(url, json.dumps({"answers": {key: "Other"}, "revision": 0}), content_type="application/json")
        self.assertEqual(conflict.status_code, 400)
        session.refresh_from_db()
        response = self.export(session)
        self.assertEqual(response.json()["job"]["status"], "complete")
        self.assertEqual(self.client.get(url).json()["session"]["fields"][0]["value"], "")

    def test_control_requires_answer_but_false_is_valid(self):
        session = self.upload("{% if attending %}Attending{% else %}Absent{% endif %}")
        field = next(f for f in session.fill_state["fields"] if f["path"] == "attending")
        self.assertTrue(field["required"])
        result = self.export(session).json()["job"]
        self.assertEqual(result["status"], "failed")
        self.assertIn("attending", result["error"])
        session.refresh_from_db()
        result = self.export(session, {field["key"]: False}).json()["job"]
        self.assertEqual(result["status"], "complete", result)
        content = self.client.get(f"/api/template-fill/jobs/{result['id']}/file/").content
        self.assertIn("Absent", " ".join(p.text for p in Document(io.BytesIO(content)).paragraphs))

    def test_loop_collection_and_user_text_not_evaluated(self):
        session = self.upload("{% for person in people %}{{ person.name }} {% endfor %}")
        self.assertEqual([f["path"] for f in session.fill_state["fields"]], ["people"])
        field = session.fill_state["fields"][0]
        result = self.export(session, {field["key"]: [{"name": "{{ untouched }}"}]}).json()["job"]
        self.assertEqual(result["status"], "complete", result)
        content = self.client.get(f"/api/template-fill/jobs/{result['id']}/file/").content
        self.assertIn("{{ untouched }}", " ".join(p.text for p in Document(io.BytesIO(content)).paragraphs))

    def test_access_control_and_private_upload_catalog(self):
        session = self.upload("{{ fields.hearing_date }}")
        other = get_user_model().objects.create_user("other", "other@example.test", "pw")
        self.client.force_login(other)
        with patch("apps.drafting.template_fill_views.user_can_access_matter", return_value=False):
            self.assertEqual(self.client.get(f"/api/template-fill/sessions/{session.pk}/").status_code, 404)
            self.assertEqual(self.client.get(f"/api/template-fill/jobs/{session.fill_jobs.first().pk}/").status_code, 404)
            self.assertEqual(self.client.get(f"/api/template-fill/uploads/{FillTemplateUpload.objects.first().pk}/file/").status_code, 404)
        other_matter = Matter.objects.create(external_id="other", client_name="Other", matter_type="Housing", jurisdiction="Ohio")
        with patch("apps.drafting.template_fill_views.matter_for_user", return_value=other_matter):
            self.assertEqual([t for t in self.client.get("/api/template-fill/?matterId=other").json()["templates"] if t["type"] == "upload"], [])
            upload = FillTemplateUpload.objects.first()
            upload.is_published = True
            upload.save()
            self.assertEqual(len([t for t in self.client.get("/api/template-fill/?matterId=other").json()["templates"] if t["type"] == "upload"]), 1)

    def test_library_generation_slot_is_manual(self):
        template = DocumentTemplate.objects.create(title="Motion", slug="fill-motion", kind="motion")
        TemplateBlock.objects.create(template=template, key="argument", label="Argument", block_type="argument", body='{{ blocks.argument.body }}', ai_latitude="generate", ai_instructions="Explain the requested relief")
        response = self.client.post("/api/template-fill/start/", json.dumps({"matterId": self.matter.external_id, "templateId": template.pk}), content_type="application/json")
        job = response.json()["job"]
        self.assertEqual(job["status"], "complete", job)
        session = DraftingSession.objects.get(pk=job["sessionId"])
        self.assertEqual(session.fill_state["fields"][0]["blockKey"], "argument")
        result = self.export(session).json()["job"]
        self.assertEqual(result["status"], "complete", result)
        content = self.client.get(f"/api/template-fill/jobs/{result['id']}/file/").content
        self.assertIn("[Enter Argument]", " ".join(p.text for p in Document(io.BytesIO(content)).paragraphs))
        self.assertEqual(self.client.post(f"/api/drafting-sessions/{session.pk}/draft/", "{}", content_type="application/json").status_code, 400)

    def test_context_names_each_blank_and_matches_whole_paths(self):
        document = Document()
        for text in ["I, {{ fields.i }}, depose and say:", "I am the {{ role }} in {{ caption }}.",
                     "{{ fields.insert_case_facts }}", "Further affiant sayeth naught.", "{{ fields.signature }}", "NOTARY PUBLIC"]:
            document.add_paragraph(text)
        output = io.BytesIO()
        document.save(output)
        fields = {field["key"]: field for field in inspect_docx(output.getvalue())}
        # "fields.i" is a prefix of "fields.insert_case_facts"; a substring
        # search once gave a field the wrong sentence.
        self.assertEqual(fields["fields['i']"]["context"], [{"text": "I, "}, {"fields": ["fields['i']"]}, {"text": ", depose and say:"}])
        self.assertEqual(fields["caption"]["context"], [{"text": "I am the "}, {"fields": ["role"]}, {"text": " in "}, {"fields": ["caption"]}, {"text": "."}])
        self.assertEqual(fields["fields['signature']"]["context"], [{"text": "Further affiant sayeth naught."}, {"break": True}, {"fields": ["fields['signature']"]}, {"break": True}, {"text": "NOTARY PUBLIC"}])

    def test_generation_slot_the_document_never_prints_is_not_offered(self):
        template = DocumentTemplate.objects.create(title="Affidavit", slug="fill-affidavit", kind="affidavit")
        TemplateBlock.objects.create(template=template, key="facts", label="Facts", block_type="facts", body='{{ blocks["facts"]["body"] }}', ai_latitude="generate", ai_instructions=["Insert case specific facts"])
        TemplateBlock.objects.create(template=template, key="state-of-ohio", label="State Of Ohio", block_type="caption", body="STATE OF OHIO", ai_latitude="generate")
        response = self.client.post("/api/template-fill/start/", json.dumps({"matterId": self.matter.external_id, "templateId": template.pk}), content_type="application/json")
        session = DraftingSession.objects.get(pk=response.json()["job"]["sessionId"])
        manual = [field for field in session.fill_state["fields"] if field.get("blockKey")]
        self.assertEqual([field["blockKey"] for field in manual], ["facts"])
        self.assertEqual(manual[0]["reason"], "Insert case specific facts")

    def test_preview_marks_values_and_prompts_and_saves_nothing(self):
        session = self.upload("Dear {{ fields.name }}, hearing {{ fields.hearing }}.{% if fields.attend %} You will attend.{% endif %}")
        revision = session.fill_state["revision"]
        response = self.client.post(f"/api/template-fill/sessions/{session.pk}/preview/", json.dumps({"answers": {"fields['name']": "Alex"}}), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        preview = response.json()["preview"]
        segments = [segment for block in preview["body"] for segment in block.get("segments", [])]
        self.assertIn({"filled": "Alex", "key": "fields['name']", "label": "name"}, segments)
        self.assertIn({"prompt": "hearing", "key": "fields['hearing']"}, segments)
        text = "".join(segment.get("text", "") for segment in segments)
        self.assertNotIn("You will attend", text)
        # An unanswered condition is previewed as off and reported; the export
        # still refuses it.
        self.assertEqual(preview["assumedOff"], ["attend"])
        session.refresh_from_db()
        self.assertEqual(session.fill_state["revision"], revision)
        self.assertEqual(session.fill_state["answers"], {})
        answered = self.client.post(f"/api/template-fill/sessions/{session.pk}/preview/", json.dumps({"answers": {"fields['attend']": True}}), content_type="application/json").json()["preview"]
        self.assertIn("You will attend", "".join(segment.get("text", "") for block in answered["body"] for segment in block["segments"]))
        self.assertEqual(answered["assumedOff"], [])
        bad = self.client.post(f"/api/template-fill/sessions/{session.pk}/preview/", json.dumps({"answers": {"unknown": "x"}}), content_type="application/json")
        self.assertEqual(bad.status_code, 400)

    def test_mapping_precedence_disabled_missing_and_seed_preserves_edits(self):
        seed_mappings()
        base = TemplateFieldMapping.objects.get(field=canonical_path("clients[0].name.first"))
        base.source_path = "custom.first"
        base.save()
        seed_mappings()
        base.refresh_from_db()
        self.assertEqual(base.source_path, "custom.first")
        specific = TemplateFieldMapping.objects.create(organization=base.organization, field=base.field, template_slug="specific", source_path="special", enabled=False)
        self.assertEqual(mapping_index("specific")[base.field], specific)
        with self.assertRaisesRegex(ValueError, "disabled"):
            mapped_value(specific, {"special": "value"})
        base.source_path = "count"
        self.assertEqual(mapped_value(base, {"count": 0}), 0)
        self.assertIs(mapped_value(base, {"count": False}), False)
        self.assertEqual(mapped_value(base, {"count": {"text_value": "N/A", "raw_value": 0}}), 0)
        with self.assertRaises(ValueError):
            mapped_value(base, {})

    def test_preparation_rejects_calls_private_paths_and_unsafe_templates(self):
        for text in ('{{ clients[0].name.full() }}', '{{ person.__class__ }}', '{% include "other" %}'):
            with self.subTest(text=text), self.assertRaises(ValueError):
                prepare_upload(docx_bytes(text))

    def test_literal_path_parser(self):
        self.assertEqual(read_path({"custom": {"hearing date": ["today"]}}, 'custom["hearing date"][0]'), "today")
        for path in ("__class__", "x[-1]", "x[10000]", "x()", "x[y]", "x.__dict__"):
            with self.subTest(path=path), self.assertRaises(ValueError):
                canonical_path(path)

    def test_background_start_returns_pending_without_processing(self):
        with override_settings(TEMPLATE_FILL_BACKGROUND=True), patch("apps.drafting.template_fill_jobs.threading.Thread") as thread:
            with self.captureOnCommitCallbacks(execute=True):
                response = self.client.post("/api/template-fill/start/", {"matterId": self.matter.external_id, "file": SimpleUploadedFile("example.docx", docx_bytes("{{ fields.hearing }}"))})
            self.assertEqual(response.status_code, 202)
            self.assertEqual(response.json()["job"]["status"], "pending")
            thread.return_value.start.assert_called_once()

    def test_upload_converts_author_marked_blanks_and_keeps_formatting(self):
        content, report = prepare_upload(docx_bytes("Hearing: [Hearing date]. Rent: ______"))
        self.assertGreaterEqual(len(report["fields"]), 2)
        state = {"context": {}, "answers": {}, "fields": report["fields"]}
        result = render_fill(content, state)
        document = Document(io.BytesIO(result))
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertIn("Maintained wording", document.paragraphs[0].text)
        self.assertIn("[Enter hearing date]", document.paragraphs[0].text)

    def test_upload_either_or_defaults_to_first_and_exposes_choice(self):
        document = Document()
        document.add_paragraph("Service", style="Heading 1")
        document.add_paragraph("I served by email.")
        document.add_paragraph("[OR]")
        document.add_paragraph("I served by regular mail.")
        output = io.BytesIO()
        document.save(output)
        response = self.client.post("/api/template-fill/start/", {"matterId": self.matter.external_id, "file": SimpleUploadedFile("service.docx", output.getvalue())})
        job = response.json()["job"]
        self.assertEqual(job["status"], "complete", job)
        session = DraftingSession.objects.get(pk=job["sessionId"])
        choices = [f for f in session.fill_state["fields"] if f.get("choices")]
        self.assertEqual(len(choices), 1)
        result = self.export(session).json()["job"]
        self.assertEqual(result["status"], "complete", result)
        content = self.client.get(f"/api/template-fill/jobs/{result['id']}/file/").content
        text = "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
        self.assertIn("served by email", text)
        self.assertNotIn("regular mail", text)

    def test_unreviewed_template_cannot_be_published_by_admin_action(self):
        from django.contrib.admin.sites import AdminSite
        from apps.templates_app.admin import FillTemplateUploadAdmin
        from django.test import RequestFactory
        self.upload("Blank {{ fields.hearing_date }}")
        admin = FillTemplateUploadAdmin(FillTemplateUpload, AdminSite())
        request = RequestFactory().post("/")
        request.user = self.user
        with patch.object(admin, "message_user"):
            admin.publish_reviewed(request, FillTemplateUpload.objects.all())
            self.assertFalse(FillTemplateUpload.objects.first().is_published)
            FillTemplateUpload.objects.update(reviewed_for_sharing=True)
            admin.publish_reviewed(request, FillTemplateUpload.objects.all())
            self.assertTrue(FillTemplateUpload.objects.first().is_published)

    def test_delivery_failure_does_not_lose_export(self):
        from apps.matters.models import LegalServerDelivery
        session = self.upload("{{ fields.hearing_date }}")
        failed = LegalServerDelivery.objects.create(matter=self.matter, kind="document", origin="template_fill", status="failed", reason="Synthetic failure")
        with patch("apps.drafting.template_fill.save_document", return_value=failed):
            result = self.export(session).json()["job"]
        self.assertEqual(result["status"], "complete", result)
        self.assertEqual(result["result"]["delivery"]["status"], "failed")
        self.assertEqual(self.client.get(f"/api/template-fill/jobs/{result['id']}/file/").status_code, 200)

    def test_highlighting_keeps_line_breaks_and_tabs_in_order(self):
        content = docx_bytes("Before\n{{ fields.missing }}\tAfter")
        fields = inspect_docx(content)
        result = render_fill(content, {"context": {}, "fields": fields, "answers": {}})
        self.assertEqual(Document(io.BytesIO(result)).paragraphs[0].text, "Maintained wording: Before\n[Enter missing]\tAfter")

    def test_missing_values_survive_formatting_filters(self):
        content = docx_bytes('{{ fields.hearing_date | truncate(5) | upper }}')
        result = render_fill(content, {"context": {}, "fields": inspect_docx(content), "answers": {}})
        self.assertIn("[Enter hearing date]", Document(io.BytesIO(result)).paragraphs[0].text)

    def test_supported_helper_is_not_offered_as_a_field(self):
        fields = inspect_docx(docx_bytes('{{ possessive(client_name) }}'))
        self.assertEqual([field["path"] for field in fields], ["client_name"])

    def test_first_page_header_and_nested_table_blanks(self):
        document = Document()
        document.sections[0].different_first_page_header_footer = True
        document.sections[0].first_page_header.paragraphs[0].text = "[Hearing date]"
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "[Service date]"
        stream = io.BytesIO()
        document.save(stream)
        content, report = prepare_upload(stream.getvalue())
        self.assertEqual(len(report["fields"]), 2)
        rendered = render_fill(content, {"context": {}, "fields": report["fields"], "answers": {}})
        result = Document(io.BytesIO(rendered))
        self.assertIn("[Enter hearing date]", result.sections[0].first_page_header.paragraphs[0].text)
        self.assertIn("[Enter service date]", result.tables[0].cell(0, 0).tables[0].cell(0, 0).text)

    def test_optional_sections_require_a_human_selection(self):
        template = DocumentTemplate.objects.create(title="Optional motion", slug="optional-fill-motion", kind="motion")
        TemplateBlock.objects.create(template=template, key="optional", label="Optional claim", block_type="argument", body='Optional claim. {% if details %}{{ details }}{% endif %}', required=False)
        response = self.client.post("/api/template-fill/start/", json.dumps({"matterId": self.matter.external_id, "templateId": template.pk}), content_type="application/json")
        job = response.json()["job"]
        self.assertEqual(job["status"], "complete", job)
        session = DraftingSession.objects.get(pk=job["sessionId"])
        choice = next(f for f in session.fill_state["fields"] if f["kind"] == "boolean")
        self.assertIs(choice["value"], False)
        result = self.export(session).json()["job"]
        self.assertEqual(result["status"], "complete", result)
        content = self.client.get(f"/api/template-fill/jobs/{result['id']}/file/").content
        self.assertNotIn("Optional claim.", "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs))
        session.refresh_from_db()
        details = next(f for f in session.fill_state["fields"] if f["path"] == "details")
        result = self.export(session, {choice["key"]: True, details["key"]: "Facts chosen by author"}).json()["job"]
        self.assertEqual(result["status"], "complete", result)
        content = self.client.get(f"/api/template-fill/jobs/{result['id']}/file/").content
        self.assertIn("Facts chosen by author", "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs))

    def test_active_export_does_not_discard_new_answers(self):
        session = self.upload("{{ fields.hearing_date }}")
        field = session.fill_state["fields"][0]
        with override_settings(TEMPLATE_FILL_BACKGROUND=True), patch("apps.drafting.template_fill_jobs.threading.Thread"):
            response = self.export(session, {field["key"]: "First answer"})
            self.assertEqual(response.status_code, 202)
            session.refresh_from_db()
            response = self.export(session, {field["key"]: "New answer"})
            self.assertEqual(response.status_code, 409)
            session.refresh_from_db()
            self.assertEqual(session.fill_state["answers"][field["key"]], "First answer")

    def test_library_composition_preserves_first_document_letterhead(self):
        from pathlib import Path
        from apps.templates_app.fill_templates import library_docx
        template = DocumentTemplate.objects.create(title="Letterhead", slug="fill-letterhead", kind="motion")
        TemplateBlock.objects.create(template=template, key="body", label="Body", block_type="facts", body="Maintained wording")
        source = Document()
        source.sections[0].header.paragraphs[0].text = "Neutral organization letterhead"
        source.sections[0].left_margin = 1234567
        source.add_paragraph("Maintained wording")
        path = Path(self.directory.name) / "letterhead.docx"
        source.save(path)
        with patch("apps.templates_app.fill_templates.block_template_path", return_value=path):
            result = Document(io.BytesIO(library_docx(template)))
        self.assertEqual(result.sections[0].header.paragraphs[0].text, "Neutral organization letterhead")
        self.assertEqual(result.sections[0].left_margin, source.sections[0].left_margin)
